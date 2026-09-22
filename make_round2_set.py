"""
make_round2_set.py  --  builds the round-2 annotation package (Corrections spec
Part A, A1 + A2) into annotation_round2/.

Reuses make_iaa_set.py's machinery: the blind column set, the Excel three-option
dropdown on `verdict`, stratified fixed-seed sampling, and source_question /
full_response pulled from the intact source file exactly as round 1 did.

Produces:
  annotation_round2/calibration_10.xlsx          (A1) 10 claims, 5 en + 5 hi, mixed verdicts
  annotation_round2/round2_annotator_A.xlsx       (A2) 50 claims, 30 en + 20 hi
  annotation_round2/round2_annotator_B.xlsx       (A2) identical copy
  annotation_round2/round2_key_DO_NOT_SHARE.csv   (A2) private: claim_id, ai_verdict, gold_verdict

Every set is BLIND (no ai_verdict / ai_reasoning / signal scores / gold_verdict).
Excludes the 56 round-1 IAA claims and every claim flagged `broken` in
claim_quality.csv; A2 additionally excludes the A1 calibration claims and is
drawn only from gold_source == "ai_agreed". Fixed seed -> reproducible.

Run:  .venv/Scripts/python.exe make_round2_set.py
"""
from __future__ import annotations
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

GOLD = "gold_final.csv"
SRC = "final annotator file(Sheet1) (1).csv"   # intact Original Query / Full Response Context
QUALITY = "claim_quality.csv"
IAA_KEY = Path("iaa") / "iaa_key_DO_NOT_SHARE.csv"   # the 56 round-1 claims to exclude
OUTDIR = Path("annotation_round2")
SEED = 42

CAL_PER_LANG = 5                       # A1: 5 en + 5 hi = 10
R2_COUNTS = {"en": 30, "hi": 20}       # A2: 30 en + 20 hi = 50
VERDICTS = ["Verified", "Hallucinated", "Uncertain"]

# blind columns, in order (A1/A2 spec)
COLS = ["claim_id", "language", "source_question", "claim_text",
        "full_response", "verdict", "evidence_url", "notes"]
COLW = {"claim_id": 34, "language": 8, "source_question": 50, "claim_text": 60,
        "full_response": 60, "verdict": 16, "evidence_url": 30, "notes": 28}
WRAP = {"source_question", "claim_text", "full_response", "notes"}


def has_dev(v):
    return any("ऀ" <= ch <= "ॿ" for ch in str(v))


def write_xlsx(df: pd.DataFrame, path: Path, sheet: str) -> None:
    """Blind sheet with a frozen header, wrapped text, and a Verified/
    Hallucinated/Uncertain dropdown on `verdict` -- same as make_iaa_set.py."""
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name=sheet)
        ws = w.sheets[sheet]
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = COLW.get(col, 18)
            ws.cell(row=1, column=i).font = Font(bold=True)
            if col in WRAP:
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")
        vcol = list(df.columns).index("verdict") + 1
        letter = ws.cell(row=1, column=vcol).column_letter
        dv = DataValidation(type="list", formula1='"Verified,Hallucinated,Uncertain"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{len(df) + 1}")


def blind_frame(sample: pd.DataFrame, src: pd.DataFrame) -> pd.DataFrame:
    q = src["Original Query"].reindex(sample["claim_id"]).values
    resp = src["Full Response Context"].reindex(sample["claim_id"]).values
    return pd.DataFrame({
        "claim_id": sample["claim_id"].values,
        "language": sample["language"].values,
        "source_question": q,
        "claim_text": sample["claim_text"].values,
        "full_response": resp,
        "verdict": "",        # fill: Verified / Hallucinated / Uncertain
        "evidence_url": "",   # fill: your best source
        "notes": "",          # optional
    })[COLS]


def pick_calibration(pool: pd.DataFrame, seed: int) -> pd.DataFrame:
    """5 per language, guaranteeing at least one of each gold verdict that exists
    in that language so the calibration discussion covers all three labels."""
    parts = []
    for lang in ("en", "hi"):
        sub = pool[pool.language == lang]
        seed_rows = []
        for v in VERDICTS:  # one of each verdict first
            g = sub[sub.gold_verdict == v]
            if len(g):
                seed_rows.append(g.sample(1, random_state=seed))
        picked = pd.concat(seed_rows) if seed_rows else sub.iloc[0:0]
        need = CAL_PER_LANG - len(picked)
        rest = sub[~sub.claim_id.isin(picked.claim_id)]
        if need > 0 and len(rest):
            picked = pd.concat([picked, rest.sample(min(need, len(rest)), random_state=seed)])
        parts.append(picked.head(CAL_PER_LANG))
    return pd.concat(parts).sample(frac=1, random_state=seed).reset_index(drop=True)


def pick_round2(pool: pd.DataFrame, seed: int) -> pd.DataFrame:
    parts = []
    for lang, n in R2_COUNTS.items():
        sub = pool[pool.language == lang]
        if len(sub) < n:
            raise SystemExit(f"not enough {lang} ai_agreed claims after exclusions: "
                             f"have {len(sub)}, need {n}")
        parts.append(sub.sample(n, random_state=seed))
    return pd.concat(parts).sample(frac=1, random_state=seed).reset_index(drop=True)


def _report(name: str, sample: pd.DataFrame, src_q) -> None:
    print(f"\n{name}: {len(sample)} claims")
    print(sample.groupby(["language", "gold_verdict"]).size().to_string())
    miss = int(pd.isna(src_q).sum())
    if miss:
        print(f"  WARNING: {miss} rows missing a source_question")


def main() -> None:
    OUTDIR.mkdir(exist_ok=True)
    gold = pd.read_csv(GOLD, encoding="utf-8")
    src = pd.read_csv(SRC, encoding="utf-8").set_index("Claim ID")
    quality = pd.read_csv(QUALITY, encoding="utf-8")

    broken = set(quality.loc[quality["quality_flag"] != "ok", "claim_id"])
    iaa1 = set(pd.read_csv(IAA_KEY, encoding="utf-8-sig")["claim_id"])
    base = gold[~gold["claim_id"].isin(broken) & ~gold["claim_id"].isin(iaa1)].copy()
    print(f"pool after excluding {len(broken)} broken + {len(iaa1)} round-1 IAA: {len(base)} claims")

    # A1 -- calibration (from all gold_source, mixed verdicts)
    cal = pick_calibration(base, SEED)
    cal_blind = blind_frame(cal, src)
    write_xlsx(cal_blind, OUTDIR / "calibration_10.xlsx", "calibration")
    _report("A1 calibration_10", cal, cal_blind["source_question"])

    # A2 -- round-2 (ai_agreed only, excluding the calibration claims)
    pool2 = base[(base["gold_source"] == "ai_agreed") & (~base["claim_id"].isin(cal["claim_id"]))]
    r2 = pick_round2(pool2, SEED)
    r2_blind = blind_frame(r2, src)
    for tag in ("A", "B"):
        write_xlsx(r2_blind, OUTDIR / f"round2_annotator_{tag}.xlsx", "round2")
    r2[["claim_id", "ai_verdict", "gold_verdict"]].to_csv(
        OUTDIR / "round2_key_DO_NOT_SHARE.csv", index=False, encoding="utf-8-sig")
    _report("A2 round2 (A/B identical)", r2, r2_blind["source_question"])

    hi = r2_blind[r2_blind["language"] == "hi"]
    print(f"\nHindi round-2 rows: {len(hi)}  Devanagari question intact: "
          f"{int(hi['source_question'].apply(has_dev).sum())}")
    sys.stdout.buffer.write(b"\nwrote annotation_round2/: calibration_10.xlsx, "
                            b"round2_annotator_A/B.xlsx, round2_key_DO_NOT_SHARE.csv (private)\n")


if __name__ == "__main__":
    main()
