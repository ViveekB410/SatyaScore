"""
relabel/assign.py -- Version 2 data, step 1c (relabel/PROTOCOL.md): builds the BLIND annotation package.

- calibration_20.xlsx: 20 claims (10 EN, 10 HI) for all four annotators, practice only.
- annotator_A..D.xlsx: every other claim goes to TWO annotators, rotating over the six pairs
  (AB, CD, AC, BD, AD, BC) separately per language so each annotator gets a similar EN/HI mix.
- assignment_key.csv: claim_id -> pair and claim metadata (source model, extraction method). Not given
  to annotators (it is not secret, but annotators should judge the claim, not its origin).

No system or AI verdict appears anywhere in the package.
    .venv/Scripts/python.exe relabel/assign.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "relabel" / "package"
SEED = 20261012
PAIRS = [("A", "B"), ("C", "D"), ("A", "C"), ("B", "D"), ("A", "D"), ("B", "C")]
CAL_PER_LANG = 10
LABELS = ["Verified", "Hallucinated", "Uncertain", "Not a claim"]
COLS = ["claim_id", "language", "source_question", "claim_text", "needs_context", "full_response",
        "verdict", "evidence_url", "notes"]
WIDTH = {"claim_id": 24, "language": 8, "source_question": 40, "claim_text": 60, "needs_context": 10,
         "full_response": 70, "verdict": 16, "evidence_url": 30, "notes": 28}
WRAP = {"source_question", "claim_text", "full_response", "notes"}


def write_sheet(df: pd.DataFrame, path: Path, sheet: str) -> None:
    out = df.assign(verdict="", evidence_url="", notes="")[COLS]
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        out.to_excel(w, index=False, sheet_name=sheet)
        ws = w.sheets[sheet]
        ws.freeze_panes = "A2"
        for i, col in enumerate(out.columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = WIDTH[col]
            ws.cell(row=1, column=i).font = Font(bold=True)
            if col in WRAP:
                for r in range(2, len(out) + 2):
                    ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical="top")
        letter = ws.cell(row=1, column=COLS.index("verdict") + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{len(out) + 1}")


def main() -> None:
    claims = pd.concat([pd.read_csv(ROOT / "relabel" / f, encoding="utf-8")
                        for f in ("claims_v2_original.csv", "claims_v2_new_hi.csv")], ignore_index=True)
    if claims.claim_id.duplicated().any():
        sys.exit("duplicate claim ids")
    OUT.mkdir(exist_ok=True)

    cal = pd.concat([claims[claims.language == lang].sample(CAL_PER_LANG, random_state=SEED)
                     for lang in ("en", "hi")]).sample(frac=1, random_state=SEED)
    rest = claims[~claims.claim_id.isin(cal.claim_id)].copy()
    rest["pair"] = ""
    for lang in ("en", "hi"):
        idx = rest[rest.language == lang].sample(frac=1, random_state=SEED).index
        for k, i in enumerate(idx):
            rest.at[i, "pair"] = "".join(PAIRS[k % len(PAIRS)])

    write_sheet(cal, OUT / "calibration_20.xlsx", "calibration")
    print(f"calibration_20.xlsx: {len(cal)} claims ({cal.language.value_counts().to_dict()})")
    for a in "ABCD":
        mine = rest[rest.pair.str.contains(a)].sample(frac=1, random_state=SEED + ord(a))
        write_sheet(mine, OUT / f"annotator_{a}.xlsx", f"annotator_{a}")
        print(f"annotator_{a}.xlsx: {len(mine)} claims ({mine.language.value_counts().to_dict()})")
    key = pd.concat([cal.assign(pair="ABCD", role="calibration"), rest.assign(role="evaluation")])
    key[["claim_id", "role", "pair", "set", "language", "domain", "source_llm", "source_query_id",
         "extraction_method", "needs_context"]].to_csv(OUT / "assignment_key.csv", index=False, encoding="utf-8")
    print(f"assignment_key.csv: {len(key)} claims; evaluation claims {len(rest)}, each labelled twice")


if __name__ == "__main__":
    main()
